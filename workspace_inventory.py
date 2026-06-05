#!/usr/bin/env python3
"""
workspace_inventory.py – Databricks Workspace Inventory Tool

Connects directly to a Databricks workspace via REST APIs (no migrate tool required)
and generates a comprehensive HTML report with a Summary dashboard and per-component
Detail views, plus an Excel workbook with a Summary sheet and one sheet per component.

Usage:
    python3 workspace_inventory.py \\
        --workspace-url https://<workspace>.azuredatabricks.net \\
        --token dapiXXXXXXXXXX \\
        [--output inventory_<workspace>_<timestamp>.html] \\
        [--excel-output inventory_<workspace>_<timestamp>.xlsx] \\
        [--max-scim 2000] \\
        [--no-ssl-verification]

Components inventoried:
    Users, Groups, Service Principals, Notebooks, Workspace Files,
    Jobs, All-Purpose Clusters, Instance Pools, Cluster Policies,
    SQL Warehouses, DLT Pipelines, AI/BI Dashboards, Genie AI Spaces,
    Secret Scopes, Git Repos, Model Serving Endpoints
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
import urllib3
from datetime import datetime
from typing import Any, Dict, List, Optional
from urllib.parse import urlparse

try:
    import requests
except ImportError:
    sys.exit("requests library is required: pip install requests")

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)


# ---------------------------------------------------------------------------
# Verbose logger – single shared instance, enabled via --verbose
# ---------------------------------------------------------------------------

class _VLog:
    """Lightweight verbose logger. Call vlog.enable() to activate."""

    def __init__(self):
        self._on = False

    def enable(self):
        self._on = True

    @property
    def is_on(self) -> bool:
        return self._on

    def __call__(self, msg: str, indent: int = 4):
        if self._on:
            ts = datetime.now().strftime("%H:%M:%S")
            print(f"  {' ' * indent}[{ts}] {msg}", flush=True)

    def api(self, method: str, url: str, status: Optional[int] = None,
            count: Optional[int] = None, elapsed_ms: Optional[float] = None):
        if not self._on:
            return
        parts = [f"{method}  {url}"]
        if status is not None:
            parts.append(f"→ HTTP {status}")
        if count is not None:
            parts.append(f"({count} items)")
        if elapsed_ms is not None:
            parts.append(f"[{elapsed_ms:.0f}ms]")
        self(("  ".join(parts)), indent=6)

    def section(self, msg: str):
        if self._on:
            print(f"\n  ── {msg}", flush=True)


vlog = _VLog()


# ---------------------------------------------------------------------------
# Databricks logo helper
# ---------------------------------------------------------------------------
def _db_icon_svg(size: int = 40) -> str:
    """Return the Databricks stacked-bricks icon as inline SVG (no text, transparent bg)."""
    icon_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                             "assets", "databricks_icon.svg")
    try:
        with open(icon_path) as f:
            svg = f.read()
        svg = svg.replace('width="64"', f'width="{size}"')
        svg = svg.replace('height="64"', f'height="{size}"')
        return svg
    except Exception:
        return ""


# ---------------------------------------------------------------------------
# API Client
# ---------------------------------------------------------------------------

class DatabricksClient:
    """Thin wrapper around the Databricks REST API."""

    def __init__(self, workspace_url: str, token: str, verify_ssl: bool = True):
        self.base_url = workspace_url.rstrip("/")
        self._s = requests.Session()
        self._s.headers.update({"Authorization": f"Bearer {token}"})
        self._s.verify = verify_ssl

    def get(self, path: str, params: Optional[Dict] = None) -> Dict:
        url = f"{self.base_url}/{path.lstrip('/')}"
        t0 = time.time()
        try:
            r = self._s.get(url, params=params, timeout=30)
            elapsed = (time.time() - t0) * 1000
            r.raise_for_status()
            result = r.json()
            vlog.api("GET", url, status=r.status_code, elapsed_ms=elapsed)
            return result
        except requests.HTTPError as e:
            elapsed = (time.time() - t0) * 1000
            status = getattr(e.response, "status_code", 0)
            vlog.api("GET", url, status=status, elapsed_ms=elapsed)
            vlog(f"  ⚠  HTTP error: {e}", indent=8)
            return {"_error": str(e), "_status": status}
        except Exception as e:
            elapsed = (time.time() - t0) * 1000
            vlog.api("GET", url, elapsed_ms=elapsed)
            vlog(f"  ⚠  Error: {e}", indent=8)
            return {"_error": str(e)}

    def get_paginated(self, path: str, result_key: str,
                      token_key: str = "next_page_token",
                      max_pages: int = 50) -> List[Dict]:
        """Paginate through results using a cursor token."""
        items: List[Dict] = []
        params: Dict[str, Any] = {}
        for page_num in range(max_pages):
            if page_num > 0:
                vlog(f"  page {page_num + 1}  (accumulated {len(items)} items so far)", indent=8)
            data = self.get(path, params=params)
            if "_error" in data:
                break
            batch = data.get(result_key, [])
            items.extend(batch if isinstance(batch, list) else [])
            cursor = data.get(token_key, "")
            if not cursor:
                break
            params = {"page_token": cursor}
        vlog(f"  paginated fetch done: {len(items)} total items", indent=8)
        return items

    def get_scim(self, resource: str, max_items: int = 0) -> List[Dict]:
        """Paginate SCIM resources (Users/Groups/ServicePrincipals).

        Args:
            max_items: If > 0, stop after fetching this many items (useful for
                       large workspaces with tens of thousands of users).
        """
        items: List[Dict] = []
        start = 1
        count = 500
        page  = 0
        while True:
            page += 1
            data = self.get(f"api/2.0/preview/scim/v2/{resource}",
                            params={"startIndex": start, "count": count})
            if "_error" in data:
                break
            resources = data.get("Resources", [])
            total     = data.get("totalResults", 0)
            items.extend(resources)
            cap_note = f"  (cap={max_items})" if max_items else ""
            vlog(f"  SCIM {resource} page {page}: "
                 f"startIndex={start}  got={len(resources)}  "
                 f"total={total}{cap_note}  accumulated={len(items)}",
                 indent=8)
            if max_items > 0 and len(items) >= max_items:
                items = items[:max_items]
                vlog(f"  SCIM {resource}: reached cap {max_items}, stopping", indent=8)
                break
            if start + count - 1 >= total or not resources:
                break
            start += count
        vlog(f"  SCIM {resource} complete: {len(items)} items fetched", indent=8)
        return items


# ---------------------------------------------------------------------------
# Inventory Fetchers
# ---------------------------------------------------------------------------

class WorkspaceInventory:
    """Fetches all component data from a Databricks workspace."""

    def __init__(self, client: DatabricksClient, max_scim: int = 0,
                 max_workspace_items: int = 0, max_ws_api_calls: int = 0):
        self._c = client
        self._max_scim      = max_scim
        self._max_ws        = max_workspace_items
        self._max_ws_calls  = max_ws_api_calls
        self._data: Dict[str, Any] = {}
        self.errors: List[str] = []

    # ── Helpers ────────────────────────────────────────────────────────────

    def _safe(self, name: str, fn):
        """Call fn(), store result; record errors."""
        t0 = time.time()
        try:
            result = fn()
            self._data[name] = result
            elapsed = time.time() - t0
            n = len(result) if isinstance(result, list) else (
                sum(result.values()) if isinstance(result, dict) else "?")
            vlog(f"✓  {name:<25}  {n} items  [{elapsed:.1f}s]", indent=4)
        except Exception as e:
            self._data[name] = []
            self.errors.append(f"{name}: {e}")
            vlog(f"✗  {name:<25}  ERROR: {e}", indent=4)

    # Any path segment matching these names will cause the whole subtree to be skipped.
    # This prevents recursing into git internals, JS build artifacts, etc.
    _WS_SKIP_SEGMENTS = {
        # Git internals
        ".git", "objects", "refs", "hooks", "info", "pack",
        # JS / frontend build artifacts
        "node_modules", ".next", "dist", "build", ".cache",
        # Python artifacts
        "__pycache__", ".eggs", "*.egg-info",
        # Misc
        ".ipynb_checkpoints", ".venv", "venv", "env",
    }

    def _ws_list_recursive(self, path: str = "/", max_items: int = 0,
                           max_api_calls: int = 0,
                           _counter: Optional[List[int]] = None) -> List[Dict]:
        """Recursively list workspace objects.

        Args:
            max_items:     Stop after finding this many file/notebook items.
            max_api_calls: Stop after making this many workspace/list API calls
                           (limits scan time on workspaces with deep source trees).
        Both counters are shared across recursion via _counter[items, calls].
        Skips .git internals, node_modules, and other non-notebook subtrees.
        """
        if _counter is None:
            _counter = [0, 0]  # [items_found, api_calls_made]
            vlog(f"  workspace scan starting  "
                 f"(max_items={max_items or '∞'}, "
                 f"max_api_calls={max_api_calls or '∞'})",
                 indent=8)
        items: List[Dict] = []

        if max_items > 0 and _counter[0] >= max_items:
            return items
        if max_api_calls > 0 and _counter[1] >= max_api_calls:
            return items

        # Skip if ANY path segment is a known non-content directory
        segments = set(path.strip("/").split("/"))
        if segments & self._WS_SKIP_SEGMENTS:
            vlog(f"  skipping  {path}  (non-content dir)", indent=8)
            return items

        vlog(f"  scanning  {path}  (items={_counter[0]}, calls={_counter[1]})",
             indent=8)
        _counter[1] += 1
        data = self._c.get("api/2.0/workspace/list", params={"path": path})
        if "_error" in data or "objects" not in data:
            return items

        for obj in data.get("objects", []):
            if max_items > 0 and _counter[0] >= max_items:
                vlog(f"  workspace scan: item cap {max_items} reached", indent=8)
                break
            if max_api_calls > 0 and _counter[1] >= max_api_calls:
                vlog(f"  workspace scan: API call cap {max_api_calls} reached", indent=8)
                break
            if obj.get("object_type") == "DIRECTORY":
                items.extend(self._ws_list_recursive(
                    obj["path"], max_items=max_items,
                    max_api_calls=max_api_calls, _counter=_counter))
            else:
                items.append(obj)
                _counter[0] += 1
        return items

    # ── Fetchers ───────────────────────────────────────────────────────────

    def fetch_all(self) -> Dict[str, Any]:
        ms = self._max_scim
        mw = self._max_ws
        steps = [
            ("users",              lambda: self._c.get_scim("Users", ms)),
            ("groups",             lambda: self._c.get_scim("Groups", ms)),
            ("service_principals", lambda: self._c.get_scim("ServicePrincipals", ms)),
            ("workspace_items",    lambda: self._ws_list_recursive(
                                       "/", max_items=mw,
                                       max_api_calls=self._max_ws_calls)),
            ("jobs",               lambda: self._c.get_paginated(
                                       "api/2.1/jobs/list", "jobs",
                                       token_key="next_page_token")),
            ("clusters",           lambda: self._c.get("api/2.0/clusters/list").get("clusters", [])),
            ("instance_pools",     lambda: self._c.get("api/2.0/instance-pools/list").get("instance_pools", [])),
            ("cluster_policies",   lambda: self._c.get("api/2.0/policies/clusters/list").get("policies", [])),
            ("sql_warehouses",     lambda: self._c.get("api/2.0/sql/warehouses").get("warehouses", [])),
            ("dlt_pipelines",      lambda: self._c.get_paginated(
                                       "api/2.0/pipelines", "statuses",
                                       token_key="next_page_token")),
            ("lakeview_dashboards",lambda: self._c.get_paginated(
                                       "api/2.0/lakeview/dashboards", "dashboards",
                                       token_key="next_page_token")),
            ("genie_spaces",       lambda: self._c.get("api/2.0/genie/spaces").get("spaces", [])),
            ("secret_scopes",      lambda: self._c.get("api/2.0/secrets/scopes/list").get("scopes", [])),
            ("repos",              lambda: self._c.get_paginated(
                                       "api/2.0/repos", "repos",
                                       token_key="next_page_token")),
            ("serving_endpoints",  lambda: self._c.get("api/2.0/serving-endpoints").get("endpoints", [])),
        ]
        total = len(steps)
        for idx, (name, fn) in enumerate(steps, 1):
            label = _LABELS.get(name, name.replace("_", " ").title())
            if vlog.is_on:
                print(f"\n  [{idx:>2}/{total}] {label}", flush=True)
            else:
                print(f"  [{idx:>2}/{total}] fetching {label:<30} …", end="\r", flush=True)
            self._safe(name, fn)
        if not vlog.is_on:
            print(" " * 70, end="\r")  # clear the progress line
        return self._data

    @property
    def data(self) -> Dict[str, Any]:
        return self._data


# ---------------------------------------------------------------------------
# HTML Report Generator
# ---------------------------------------------------------------------------

_ICONS = {
    "users":              ("👤", "#4f46e5"),
    "groups":             ("👥", "#7c3aed"),
    "service_principals": ("🔑", "#9333ea"),
    "workspace_items":    ("📁", "#0891b2"),
    "notebooks":          ("📓", "#0284c7"),
    "workspace_files":    ("📄", "#0369a1"),
    "jobs":               ("⚙️",  "#d97706"),
    "clusters":           ("🖥️",  "#059669"),
    "instance_pools":     ("🏊",  "#10b981"),
    "cluster_policies":   ("📋",  "#6d28d9"),
    "sql_warehouses":     ("🗄️",  "#dc2626"),
    "dlt_pipelines":      ("🔀",  "#ea580c"),
    "lakeview_dashboards":("📊",  "#be185d"),
    "genie_spaces":       ("✨",  "#7c3aed"),
    "secret_scopes":      ("🔒",  "#0f766e"),
    "repos":              ("📦",  "#1d4ed8"),
    "serving_endpoints":  ("🚀",  "#b91c1c"),
}

_LABELS = {
    "users":              "Users",
    "groups":             "Groups",
    "service_principals": "Service Principals",
    "workspace_items":    "Workspace Items",
    "notebooks":          "Notebooks",
    "workspace_files":    "Workspace Files",
    "jobs":               "Jobs",
    "clusters":           "All-Purpose Clusters",
    "instance_pools":     "Instance Pools",
    "cluster_policies":   "Cluster Policies",
    "sql_warehouses":     "SQL Warehouses",
    "dlt_pipelines":      "DLT Pipelines",
    "lakeview_dashboards":"AI/BI Dashboards",
    "genie_spaces":       "Genie AI Spaces",
    "secret_scopes":      "Secret Scopes",
    "repos":              "Git Repos",
    "serving_endpoints":  "Model Serving Endpoints",
}

# Column definitions per component: (key_in_obj, display_label, cell_formatter_name)
_COLUMNS: Dict[str, List[tuple]] = {
    "users": [
        ("userName",    "Username",      "plain"),
        ("displayName", "Display Name",  "plain"),
        ("active",      "Active",        "badge_bool"),
        ("emails",      "Email",         "first_email"),
        ("id",          "SCIM ID",       "mono"),
    ],
    "groups": [
        ("displayName",   "Group Name",    "plain"),
        ("id",            "SCIM ID",       "mono"),
        ("members",       "Members",       "count"),
        ("roles",         "Roles",         "count"),
        ("entitlements",  "Entitlements",  "list_vals"),
    ],
    "service_principals": [
        ("displayName",   "Display Name",  "plain"),
        ("applicationId", "App ID",        "mono"),
        ("active",        "Active",        "badge_bool"),
        ("id",            "SCIM ID",       "mono"),
    ],
    "workspace_items": [
        ("path",        "Path",          "path"),
        ("object_type", "Type",          "badge_type"),
        ("language",    "Language",      "badge_lang"),
        ("object_id",   "Object ID",     "mono"),
    ],
    "jobs": [
        ("settings.name",     "Job Name",      "plain"),
        ("job_id",            "Job ID",        "mono"),
        ("settings.schedule", "Schedule",      "schedule"),
        ("creator_user_name", "Creator",       "plain"),
        ("created_time",      "Created",       "epoch_ms"),
    ],
    "clusters": [
        ("cluster_name",    "Cluster Name",    "plain"),
        ("cluster_id",      "Cluster ID",      "mono"),
        ("state",           "State",           "badge_state"),
        ("cluster_source",  "Source",          "plain"),
        ("spark_version",   "Spark Version",   "plain"),
        ("node_type_id",    "Node Type",       "plain"),
        ("autotermination_minutes", "Auto-Term (min)", "plain"),
        ("creator_user_name","Creator",         "plain"),
    ],
    "instance_pools": [
        ("instance_pool_name", "Pool Name",      "plain"),
        ("instance_pool_id",   "Pool ID",        "mono"),
        ("node_type_id",       "Node Type",      "plain"),
        ("state",              "State",          "badge_state"),
        ("min_idle_instances", "Min Idle",       "plain"),
        ("max_capacity",       "Max Capacity",   "plain"),
    ],
    "cluster_policies": [
        ("name",              "Policy Name",   "plain"),
        ("policy_id",         "Policy ID",     "mono"),
        ("description",       "Description",   "trunc"),
        ("created_at_timestamp","Created",      "epoch_ms"),
    ],
    "sql_warehouses": [
        ("name",              "Warehouse Name", "plain"),
        ("id",                "ID",             "mono"),
        ("state",             "State",          "badge_state"),
        ("warehouse_type",    "Type",           "plain"),
        ("cluster_size",      "Size",           "plain"),
        ("num_clusters",      "Clusters",       "plain"),
        ("auto_stop_mins",    "Auto-Stop (min)","plain"),
        ("creator_name",      "Creator",        "plain"),
    ],
    "dlt_pipelines": [
        ("name",              "Pipeline Name",  "plain"),
        ("pipeline_id",       "Pipeline ID",    "mono"),
        ("state",             "State",          "badge_state"),
        ("cluster_label",     "Cluster",        "plain"),
        ("creator_user_name", "Creator",        "plain"),
        ("continuous",        "Continuous",     "badge_bool"),
    ],
    "lakeview_dashboards": [
        ("display_name",      "Dashboard Name", "plain"),
        ("dashboard_id",      "Dashboard ID",   "mono"),
        ("lifecycle_state",   "State",          "badge_state"),
        ("create_time",       "Created",        "iso_ts"),
        ("update_time",       "Updated",        "iso_ts"),
    ],
    "genie_spaces": [
        ("title",             "Space Name",     "plain"),
        ("space_id",          "Space ID",       "mono"),
        ("description",       "Description",    "trunc"),
        ("warehouse_id",      "Warehouse ID",   "mono"),
        ("created_timestamp", "Created",        "epoch_ms"),
    ],
    "secret_scopes": [
        ("name",              "Scope Name",     "plain"),
        ("backend_type",      "Backend",        "badge_type"),
        ("keyvault_metadata", "Key Vault",      "kv_dns"),
    ],
    "repos": [
        ("path",              "Path",           "path"),
        ("url",               "Repository URL", "url_link"),
        ("provider",          "Provider",       "badge_type"),
        ("branch",            "Branch",         "plain"),
        ("head_commit_id",    "Commit",         "short_mono"),
    ],
    "serving_endpoints": [
        ("name",              "Endpoint Name",  "plain"),
        ("state.ready",       "Ready",          "plain"),
        ("creator",           "Creator",        "plain"),
        ("creation_timestamp","Created",        "epoch_ms"),
        ("last_updated_timestamp","Updated",    "epoch_ms"),
    ],
}


def _deep_get(obj: Any, dotted_key: str) -> Any:
    """Get a value from a nested dict using dot notation."""
    for part in dotted_key.split("."):
        if isinstance(obj, dict):
            obj = obj.get(part)
        else:
            return None
    return obj


def _cell_html(value: Any, fmt: str) -> str:
    """Render a table cell value as HTML based on formatter name."""
    if value is None or value == "":
        return '<span class="na">—</span>'

    if fmt == "plain":
        return _esc(str(value))

    if fmt == "mono":
        return f'<code class="mono">{_esc(str(value))}</code>'

    if fmt == "short_mono":
        s = str(value)[:8]
        return f'<code class="mono">{_esc(s)}</code>'

    if fmt == "path":
        return f'<span class="path">{_esc(str(value))}</span>'

    if fmt == "trunc":
        s = str(value)
        if len(s) > 80:
            s = s[:77] + "…"
        return _esc(s)

    if fmt == "badge_bool":
        if value is True or str(value).lower() in ("true", "1", "yes"):
            return '<span class="badge badge-green">Yes</span>'
        return '<span class="badge badge-red">No</span>'

    if fmt == "badge_state":
        state = str(value).upper()
        color = {
            "RUNNING": "green", "ACTIVE": "green", "PUBLISHED": "green",
            "STARTED": "green", "READY": "green", "SUCCEEDED": "green",
            "STOPPED": "gray",  "TERMINATED": "gray", "IDLE": "gray",
            "STARTING": "blue", "PENDING": "blue", "RESIZING": "blue",
            "FAILED": "red",    "ERROR": "red",    "DELETED": "red",
            "DRAFT": "yellow",  "DEPLOYING": "yellow", "INITIALIZING": "yellow",
        }.get(state, "gray")
        return f'<span class="badge badge-{color}">{_esc(state)}</span>'

    if fmt == "badge_type":
        return f'<span class="badge badge-blue">{_esc(str(value))}</span>'

    if fmt == "badge_lang":
        lang = str(value).upper() if value else ""
        color = {"PYTHON": "blue", "SCALA": "red", "SQL": "green",
                 "R": "yellow", "AUTO": "gray"}.get(lang, "gray")
        return f'<span class="badge badge-{color}">{_esc(lang or "—")}</span>' if lang else '<span class="na">—</span>'

    if fmt == "count":
        n = len(value) if isinstance(value, list) else 0
        return f'<span class="count">{n}</span>'

    if fmt == "first_email":
        if isinstance(value, list) and value:
            return _esc(value[0].get("value", ""))
        return '<span class="na">—</span>'

    if fmt == "list_vals":
        if isinstance(value, list) and value:
            vals = ", ".join(str(v.get("value", v)) for v in value[:3])
            if len(value) > 3:
                vals += f" (+{len(value)-3})"
            return _esc(vals)
        return '<span class="na">—</span>'

    if fmt == "schedule":
        if isinstance(value, dict):
            cron = value.get("quartz_cron_expression", "")
            tz = value.get("timezone_id", "")
            if cron:
                return f'<span class="schedule">{_esc(cron)}<br><small>{_esc(tz)}</small></span>'
        return '<span class="na">Manual</span>'

    if fmt == "epoch_ms":
        try:
            ts = int(value) / 1000
            return datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M")
        except Exception:
            return _esc(str(value))

    if fmt == "iso_ts":
        try:
            s = str(value)[:16].replace("T", " ")
            return _esc(s)
        except Exception:
            return _esc(str(value))

    if fmt == "url_link":
        url = str(value)
        display = url.replace("https://", "").replace("http://", "")
        if len(display) > 50:
            display = display[:47] + "…"
        return f'<a href="{_esc(url)}" target="_blank">{_esc(display)}</a>'

    if fmt == "kv_dns":
        if isinstance(value, dict):
            return _esc(value.get("dns_name", ""))
        return '<span class="na">—</span>'

    return _esc(str(value))


def _esc(s: str) -> str:
    return (s.replace("&", "&amp;").replace("<", "&lt;")
             .replace(">", "&gt;").replace('"', "&quot;"))


def _build_summary_data(data: Dict[str, Any]) -> Dict[str, int]:
    """Compute per-component counts, splitting workspace_items by type."""
    counts: Dict[str, int] = {}
    ws_items = data.get("workspace_items", [])
    counts["notebooks"]       = sum(1 for x in ws_items if x.get("object_type") == "NOTEBOOK")
    counts["workspace_files"] = sum(1 for x in ws_items if x.get("object_type") == "FILE")

    for key in ["users", "groups", "service_principals",
                "jobs", "clusters", "instance_pools", "cluster_policies",
                "sql_warehouses", "dlt_pipelines", "lakeview_dashboards",
                "genie_spaces", "secret_scopes", "repos", "serving_endpoints"]:
        counts[key] = len(data.get(key, []))

    return counts


# ---------------------------------------------------------------------------
# HTML Template
# ---------------------------------------------------------------------------

def _render_html(workspace_url: str, data: Dict[str, Any],
                 errors: List[str], generated_at: str) -> str:

    counts = _build_summary_data(data)
    hostname = urlparse(workspace_url).hostname or workspace_url

    # ── Summary cards ───────────────────────────────────────────────────
    def _card(key: str, count: int) -> str:
        icon, color = _ICONS.get(key, ("📦", "#6b7280"))
        label = _LABELS.get(key, key.replace("_", " ").title())
        return f"""
        <div class="card" onclick="showTab('{key}')" data-tab="{key}">
          <div class="card-icon" style="background:{color}18;color:{color}">{icon}</div>
          <div class="card-body">
            <div class="card-count" style="color:{color}">{count}</div>
            <div class="card-label">{label}</div>
          </div>
        </div>"""

    summary_card_keys = [
        "users", "groups", "service_principals",
        "notebooks", "workspace_files",
        "jobs", "clusters", "instance_pools", "cluster_policies",
        "sql_warehouses", "dlt_pipelines",
        "lakeview_dashboards", "genie_spaces",
        "secret_scopes", "repos", "serving_endpoints",
    ]
    cards_html = "".join(_card(k, counts.get(k, 0)) for k in summary_card_keys)

    # ── Nav tabs ─────────────────────────────────────────────────────────
    def _nav_item(key: str, count: int) -> str:
        icon, color = _ICONS.get(key, ("📦", "#6b7280"))
        label = _LABELS.get(key, key.replace("_", " ").title())
        return f"""<li class="nav-item" id="nav-{key}" onclick="showTab('{key}')">
          <span class="nav-icon">{icon}</span>
          <span class="nav-label">{label}</span>
          <span class="nav-badge" style="background:{color}">{count}</span>
        </li>"""

    nav_html = '<li class="nav-item active" id="nav-summary" onclick="showTab(\'summary\')">' \
               '<span class="nav-icon">🏠</span><span class="nav-label">Summary</span></li>'
    for k in summary_card_keys:
        nav_html += _nav_item(k, counts.get(k, 0))

    # ── Detail panels ────────────────────────────────────────────────────
    def _detail_panel(key: str, items: List[Dict], cols: List[tuple]) -> str:
        icon, color = _ICONS.get(key, ("📦", "#6b7280"))
        label = _LABELS.get(key, key.replace("_", " ").title())
        n = len(items)

        # Table header
        th = "".join(f"<th onclick=\"sortTable('{key}',{i})\">{c[1]} <span class='sort-icon'>↕</span></th>"
                     for i, c in enumerate(cols))

        # Table rows – every row carries a data-idx for pagination
        rows = []
        for idx, item in enumerate(items):
            tds = "".join(
                f"<td>{_cell_html(_deep_get(item, col[0]), col[2])}</td>"
                for col in cols
            )
            rows.append(f'<tr data-idx="{idx}">{tds}</tr>')
        tbody = "\n".join(rows) if rows else \
            f'<tr><td colspan="{len(cols)}" class="empty-row">No {label.lower()} found</td></tr>'

        return f"""
  <div class="panel" id="panel-{key}">
    <div class="panel-header" style="border-left:4px solid {color}">
      <span class="panel-icon">{icon}</span>
      <h2 class="panel-title">{label}</h2>
      <span class="panel-count" style="background:{color}18;color:{color}">{n} item{'s' if n!=1 else ''}</span>
      <div class="panel-controls">
        <div class="panel-search">
          <span class="search-icon">🔍</span>
          <input type="text" id="search-{key}" placeholder="Search {label.lower()}…"
                 oninput="onSearch('{key}', this.value)">
        </div>
        <div class="page-size-wrap">
          <label for="ps-{key}">Rows per page</label>
          <input type="number" id="ps-{key}" class="page-size-input" value="25" min="1" max="1000"
                 onchange="onPageSizeChange('{key}', this.value)">
        </div>
      </div>
    </div>
    <div class="table-wrap">
      <table id="table-{key}">
        <thead><tr>{th}</tr></thead>
        <tbody id="tbody-{key}">{tbody}</tbody>
      </table>
    </div>
    <div class="pagination-bar" id="pager-{key}">
      <span class="pager-info" id="pager-info-{key}"></span>
      <div class="pager-buttons">
        <button class="pager-btn" id="btn-first-{key}"  onclick="goPage('{key}','first')"  title="First page">«</button>
        <button class="pager-btn" id="btn-prev-{key}"   onclick="goPage('{key}','prev')"   title="Previous page">‹</button>
        <span class="pager-pages" id="pager-pages-{key}"></span>
        <button class="pager-btn" id="btn-next-{key}"   onclick="goPage('{key}','next')"   title="Next page">›</button>
        <button class="pager-btn" id="btn-last-{key}"   onclick="goPage('{key}','last')"   title="Last page">»</button>
      </div>
    </div>
  </div>"""

    panels_html = ""
    for key in summary_card_keys:
        items = data.get(key)
        if items is None:
            # Split workspace_items into notebooks and files
            if key == "notebooks":
                items = [x for x in data.get("workspace_items", []) if x.get("object_type") == "NOTEBOOK"]
            elif key == "workspace_files":
                items = [x for x in data.get("workspace_items", []) if x.get("object_type") == "FILE"]
            else:
                items = []
        cols = _COLUMNS.get(key, [("path", "Path", "plain")])
        panels_html += _detail_panel(key, items, cols)

    # ── Errors section ───────────────────────────────────────────────────
    errors_html = ""
    if errors:
        items_html = "".join(f"<li>{_esc(e)}</li>" for e in errors)
        errors_html = f'<div class="errors-box"><strong>⚠ Fetch warnings:</strong><ul>{items_html}</ul></div>'

    total_items = sum(counts.values())

    # ── Full HTML ────────────────────────────────────────────────────────
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Databricks Inventory – {_esc(hostname)}</title>
<style>
  /* ── Reset & Base ── */
  *, *::before, *::after {{ box-sizing: border-box; margin: 0; padding: 0 }}
  :root {{
    --bg:       #f8fafc;
    --surface:  #ffffff;
    --sidebar:  #0f172a;
    --sidebar-hover: #1e293b;
    --sidebar-active: #1e40af;
    --text:     #1e293b;
    --text-muted: #64748b;
    --border:   #e2e8f0;
    --radius:   10px;
    --shadow:   0 1px 3px 0 rgb(0 0 0/0.1), 0 1px 2px -1px rgb(0 0 0/0.1);
    --shadow-lg:0 10px 15px -3px rgb(0 0 0/0.1), 0 4px 6px -4px rgb(0 0 0/0.1);
  }}
  html, body {{ height:100%; font-family: -apple-system,BlinkMacSystemFont,'Segoe UI','Helvetica Neue',Arial,sans-serif; background:var(--bg); color:var(--text); font-size:14px; line-height:1.5 }}
  a {{ color:#2563eb; text-decoration:none }}
  a:hover {{ text-decoration:underline }}

  /* ── Layout ── */
  .app {{ display:flex; height:100vh; overflow:hidden }}

  /* ── Sidebar ── */
  .sidebar {{
    width:240px; min-width:240px; background:var(--sidebar); color:#e2e8f0;
    display:flex; flex-direction:column; overflow-y:auto; flex-shrink:0;
  }}
  .sidebar-brand {{
    padding:20px 16px 12px; border-bottom:1px solid #1e293b;
  }}
  .sidebar-logo {{ display:flex; align-items:center; margin-bottom:12px }}
  .sidebar-logo svg {{ flex-shrink:0 }}
  .sidebar-brand h1 {{ font-size:13px; font-weight:700; color:#f1f5f9; letter-spacing:.5px; text-transform:uppercase }}
  .sidebar-brand p {{ font-size:11px; color:#64748b; margin-top:2px; word-break:break-all }}
  .sidebar-brand .ts {{ font-size:10px; color:#475569; margin-top:6px }}

  nav ul {{ list-style:none; padding:8px 0 }}
  .nav-item {{
    display:flex; align-items:center; gap:8px; padding:7px 14px;
    cursor:pointer; border-radius:6px; margin:1px 6px; transition:background .15s;
    font-size:12.5px; color:#94a3b8;
  }}
  .nav-item:hover {{ background:var(--sidebar-hover); color:#e2e8f0 }}
  .nav-item.active {{ background:var(--sidebar-active); color:#fff; font-weight:600 }}
  .nav-icon {{ font-size:15px; flex-shrink:0 }}
  .nav-label {{ flex:1; white-space:nowrap; overflow:hidden; text-overflow:ellipsis }}
  .nav-badge {{
    font-size:10px; font-weight:700; color:#fff; padding:1px 6px;
    border-radius:20px; flex-shrink:0; min-width:20px; text-align:center;
  }}

  /* ── Main Content ── */
  .main {{ flex:1; overflow-y:auto; padding:24px }}

  /* ── Summary Panel ── */
  #panel-summary {{ display:block }}
  .summary-header {{ margin-bottom:24px }}
  .summary-header h2 {{ font-size:22px; font-weight:700; color:var(--text) }}
  .summary-header p {{ color:var(--text-muted); margin-top:4px }}
  .summary-stats {{ display:flex; gap:16px; margin-top:12px; flex-wrap:wrap }}
  .stat-pill {{
    background:var(--surface); border:1px solid var(--border); border-radius:20px;
    padding:4px 14px; font-size:12px; color:var(--text-muted);
  }}
  .stat-pill strong {{ color:var(--text) }}

  .cards-grid {{
    display:grid; grid-template-columns:repeat(auto-fill, minmax(168px,1fr)); gap:12px;
    margin-bottom:24px;
  }}
  .card {{
    background:var(--surface); border:1px solid var(--border); border-radius:var(--radius);
    padding:16px; cursor:pointer; transition:all .2s; display:flex; align-items:center; gap:14px;
    box-shadow:var(--shadow);
  }}
  .card:hover {{ transform:translateY(-2px); box-shadow:var(--shadow-lg); border-color:#93c5fd }}
  .card-icon {{ font-size:24px; width:44px; height:44px; border-radius:10px; display:flex; align-items:center; justify-content:center; flex-shrink:0 }}
  .card-count {{ font-size:24px; font-weight:800; line-height:1 }}
  .card-label {{ font-size:11.5px; color:var(--text-muted); margin-top:2px; line-height:1.3 }}

  /* ── Error box ── */
  .errors-box {{
    background:#fef2f2; border:1px solid #fecaca; border-radius:var(--radius);
    padding:12px 16px; color:#b91c1c; font-size:12.5px; margin-top:16px;
  }}
  .errors-box ul {{ margin-top:6px; padding-left:16px }}

  /* ── Detail panels ── */
  .panel {{ display:none }}
  .panel.active {{ display:block }}
  .panel-header {{
    display:flex; align-items:center; gap:12px; padding:16px;
    background:var(--surface); border:1px solid var(--border);
    border-radius:var(--radius) var(--radius) 0 0; margin-bottom:0;
    flex-wrap:wrap;
  }}
  .panel-icon {{ font-size:22px }}
  .panel-title {{ font-size:18px; font-weight:700; flex-shrink:0 }}
  .panel-count {{
    font-size:12px; font-weight:700; padding:3px 10px; border-radius:20px; flex-shrink:0;
  }}
  .panel-controls {{ margin-left:auto; display:flex; align-items:center; gap:12px; flex-wrap:wrap }}
  .panel-search {{ position:relative; display:flex; align-items:center }}
  .search-icon {{ position:absolute; left:9px; font-size:13px; pointer-events:none }}
  .panel-search input {{
    border:1px solid var(--border); border-radius:6px; padding:6px 12px 6px 30px;
    font-size:13px; outline:none; width:220px; background:var(--bg);
  }}
  .panel-search input:focus {{ border-color:#93c5fd; background:#fff }}
  .page-size-wrap {{ display:flex; align-items:center; gap:6px; white-space:nowrap; color:var(--text-muted); font-size:12px }}
  .page-size-input {{
    width:62px; border:1px solid var(--border); border-radius:6px;
    padding:5px 8px; font-size:13px; text-align:center; outline:none; background:var(--bg);
  }}
  .page-size-input:focus {{ border-color:#93c5fd; background:#fff }}

  .table-wrap {{
    overflow-x:auto; background:var(--surface);
    border:1px solid var(--border); border-top:none; border-bottom:none;
  }}
  /* ── Pagination bar ── */
  .pagination-bar {{
    display:flex; align-items:center; justify-content:space-between;
    padding:10px 16px; background:var(--surface);
    border:1px solid var(--border); border-top:1px solid #f1f5f9;
    border-radius:0 0 var(--radius) var(--radius); flex-wrap:wrap; gap:8px;
  }}
  .pager-info {{ font-size:12.5px; color:var(--text-muted) }}
  .pager-info strong {{ color:var(--text) }}
  .pager-buttons {{ display:flex; align-items:center; gap:4px }}
  .pager-btn {{
    border:1px solid var(--border); background:var(--surface); color:var(--text);
    border-radius:6px; width:32px; height:32px; font-size:15px; cursor:pointer;
    display:flex; align-items:center; justify-content:center; transition:all .15s;
    padding:0;
  }}
  .pager-btn:hover:not(:disabled) {{ background:#dbeafe; border-color:#93c5fd; color:#1d4ed8 }}
  .pager-btn:disabled {{ opacity:.35; cursor:default }}
  .pager-pages {{ display:flex; gap:3px; align-items:center }}
  .page-num {{
    border:1px solid var(--border); background:var(--surface); color:var(--text-muted);
    border-radius:6px; min-width:32px; height:32px; font-size:12.5px; cursor:pointer;
    display:flex; align-items:center; justify-content:center; padding:0 6px;
    transition:all .15s; font-weight:500;
  }}
  .page-num:hover {{ background:#eff6ff; border-color:#93c5fd }}
  .page-num.current {{ background:#1d4ed8; border-color:#1d4ed8; color:#fff; font-weight:700 }}
  .page-ellipsis {{ color:var(--text-muted); padding:0 4px; font-size:12px; user-select:none }}
  table {{ width:100%; border-collapse:collapse; font-size:13px }}
  thead {{ position:sticky; top:0; z-index:2 }}
  th {{
    background:#f1f5f9; color:var(--text-muted); font-weight:600;
    padding:10px 14px; text-align:left; white-space:nowrap;
    border-bottom:2px solid var(--border); cursor:pointer; user-select:none;
  }}
  th:hover {{ background:#e2e8f0; color:var(--text) }}
  .sort-icon {{ opacity:.4; font-size:10px }}
  td {{ padding:9px 14px; border-bottom:1px solid #f1f5f9; vertical-align:top }}
  tr:last-child td {{ border-bottom:none }}
  tr:hover td {{ background:#f8fafc }}
  .empty-row {{ text-align:center; color:var(--text-muted); padding:32px; font-style:italic }}

  /* ── Badges ── */
  .badge {{
    display:inline-block; font-size:11px; font-weight:700; padding:2px 7px;
    border-radius:4px; white-space:nowrap; letter-spacing:.3px;
  }}
  .badge-green  {{ background:#dcfce7; color:#15803d }}
  .badge-red    {{ background:#fee2e2; color:#b91c1c }}
  .badge-blue   {{ background:#dbeafe; color:#1d4ed8 }}
  .badge-yellow {{ background:#fef9c3; color:#854d0e }}
  .badge-gray   {{ background:#f1f5f9; color:#475569 }}
  .badge-purple {{ background:#f3e8ff; color:#6d28d9 }}

  /* ── Misc ── */
  .na {{ color:#cbd5e1; font-style:italic }}
  .mono {{ font-family:'SF Mono','Fira Code',monospace; font-size:11.5px; color:#0f766e; background:#f0fdfa; padding:1px 5px; border-radius:3px }}
  .path {{ font-family:'SF Mono','Fira Code',monospace; font-size:11.5px; color:#7c3aed }}
  .count {{ font-weight:700; color:#1d4ed8 }}
  .schedule {{ font-size:11.5px; font-family:monospace }}
  small {{ color:var(--text-muted); font-size:11px }}
</style>
</head>
<body>
<div class="app">

<!-- ── Sidebar ─────────────────────────────────────────────────────── -->
<aside class="sidebar">
  <div class="sidebar-brand">
    <div class="sidebar-logo">
      {_db_icon_svg(48)}
    </div>
    <h1>Workspace Inventory</h1>
    <p>{_esc(hostname)}</p>
    <div class="ts">Generated {_esc(generated_at)}</div>
  </div>
  <nav><ul id="nav-list">{nav_html}</ul></nav>
</aside>

<!-- ── Main ─────────────────────────────────────────────────────────── -->
<main class="main" id="main">

  <!-- Summary Panel -->
  <div id="panel-summary" class="panel active">
    <div class="summary-header">
      <h2>Workspace Inventory</h2>
      <p>Complete inventory of all resources in <strong>{_esc(hostname)}</strong></p>
      <div class="summary-stats">
        <span class="stat-pill"><strong>{total_items}</strong> total resources</span>
        <span class="stat-pill"><strong>{len(summary_card_keys)}</strong> component types</span>
        <span class="stat-pill">Snapshot: <strong>{_esc(generated_at)}</strong></span>
      </div>
    </div>
    <div class="cards-grid">{cards_html}</div>
    {errors_html}
  </div>

  <!-- Detail Panels -->
  {panels_html}

</main>
</div>

<script>
// ─────────────────────────────────────────────────────────────────────────────
// State: per-table pagination + filter state
// ─────────────────────────────────────────────────────────────────────────────
const _state = {{}};  // keyed by tableId

function _getState(id) {{
  if (!_state[id]) _state[id] = {{ page: 1, pageSize: 25, query: '' }};
  return _state[id];
}}

// ─────────────────────────────────────────────────────────────────────────────
// Core render: applies filter + pagination to a table
// ─────────────────────────────────────────────────────────────────────────────
function _renderTable(id) {{
  const st   = _getState(id);
  const tbody = document.getElementById('tbody-' + id);
  if (!tbody) return;

  const allRows = Array.from(tbody.querySelectorAll('tr'));
  if (!allRows.length) return;

  // 1. Filter
  const q = st.query.toLowerCase().trim();
  const visible = allRows.filter(r => !q || r.textContent.toLowerCase().includes(q));

  // 2. Clamp page
  const ps       = Math.max(1, st.pageSize);
  const totalPgs = Math.max(1, Math.ceil(visible.length / ps));
  if (st.page > totalPgs) st.page = totalPgs;
  if (st.page < 1)        st.page = 1;

  // 3. Show/hide rows
  const start = (st.page - 1) * ps;
  const end   = start + ps;
  allRows.forEach(r => {{ r.style.display = 'none'; }});
  visible.forEach((r, i) => {{ r.style.display = (i >= start && i < end) ? '' : 'none'; }});

  // 4. Update info label
  const infoEl = document.getElementById('pager-info-' + id);
  if (infoEl) {{
    const from = visible.length ? start + 1 : 0;
    const to   = Math.min(end, visible.length);
    const filterNote = q ? ` (filtered from ${{allRows.length}})` : '';
    infoEl.innerHTML = visible.length
      ? `Showing <strong>${{from}}–${{to}}</strong> of <strong>${{visible.length}}</strong>${{filterNote}}`
      : `<span style="color:#b91c1c">No results match your search</span>`;
  }}

  // 5. Render page buttons
  _renderPageButtons(id, st.page, totalPgs);

  // 6. Enable/disable nav buttons
  ['first','prev'].forEach(d => {{
    const b = document.getElementById(`btn-${{d}}-${{id}}`);
    if (b) b.disabled = st.page <= 1;
  }});
  ['next','last'].forEach(d => {{
    const b = document.getElementById(`btn-${{d}}-${{id}}`);
    if (b) b.disabled = st.page >= totalPgs;
  }});
}}

function _renderPageButtons(id, current, total) {{
  const container = document.getElementById('pager-pages-' + id);
  if (!container) return;
  container.innerHTML = '';

  // Compute which page numbers to show (window around current + first/last)
  const pages = new Set();
  pages.add(1);
  pages.add(total);
  for (let p = Math.max(1, current - 2); p <= Math.min(total, current + 2); p++) pages.add(p);
  const sorted = [...pages].sort((a,b) => a-b);

  let prev = 0;
  sorted.forEach(p => {{
    if (p - prev > 1) {{
      const el = document.createElement('span');
      el.className = 'page-ellipsis';
      el.textContent = '…';
      container.appendChild(el);
    }}
    const btn = document.createElement('button');
    btn.className = 'page-num' + (p === current ? ' current' : '');
    btn.textContent = p;
    btn.onclick = () => goPage(id, p);
    container.appendChild(btn);
    prev = p;
  }});
}}

// ─────────────────────────────────────────────────────────────────────────────
// Public API called by HTML event handlers
// ─────────────────────────────────────────────────────────────────────────────
function goPage(id, target) {{
  const st = _getState(id);
  const tbody = document.getElementById('tbody-' + id);
  const allRows = tbody ? Array.from(tbody.querySelectorAll('tr')) : [];
  const q = st.query.toLowerCase().trim();
  const visible = allRows.filter(r => !q || r.textContent.toLowerCase().includes(q));
  const totalPgs = Math.max(1, Math.ceil(visible.length / st.pageSize));

  if      (target === 'first') st.page = 1;
  else if (target === 'prev')  st.page = Math.max(1, st.page - 1);
  else if (target === 'next')  st.page = Math.min(totalPgs, st.page + 1);
  else if (target === 'last')  st.page = totalPgs;
  else                          st.page = Number(target);

  _renderTable(id);
  // scroll table into view
  document.getElementById('table-' + id)?.scrollIntoView({{ block: 'nearest' }});
}}

function onSearch(id, query) {{
  const st = _getState(id);
  st.query = query;
  st.page  = 1;
  _renderTable(id);
}}

function onPageSizeChange(id, val) {{
  const n = parseInt(val, 10);
  if (!n || n < 1) return;
  const st = _getState(id);
  st.pageSize = n;
  st.page = 1;
  _renderTable(id);
}}

// ─────────────────────────────────────────────────────────────────────────────
// Sort (re-renders pagination after sorting)
// ─────────────────────────────────────────────────────────────────────────────
const _sortState = {{}};
function sortTable(tableId, colIdx) {{
  const table = document.getElementById('table-' + tableId);
  const tbody = table.querySelector('tbody');
  const rows  = Array.from(tbody.querySelectorAll('tr'));
  const key   = tableId + '_' + colIdx;
  const asc   = _sortState[key] !== true;
  _sortState[key] = asc;

  rows.sort((a, b) => {{
    const ta = a.cells[colIdx]?.textContent.trim() ?? '';
    const tb = b.cells[colIdx]?.textContent.trim() ?? '';
    const na = parseFloat(ta), nb = parseFloat(tb);
    if (!isNaN(na) && !isNaN(nb)) return asc ? na - nb : nb - na;
    return asc ? ta.localeCompare(tb) : tb.localeCompare(ta);
  }});
  rows.forEach(r => tbody.appendChild(r));

  // Update sort icons
  table.querySelectorAll('th .sort-icon').forEach((ic, i) => {{
    ic.textContent = i === colIdx ? (asc ? '↑' : '↓') : '↕';
    ic.style.opacity = i === colIdx ? '1' : '0.4';
  }});

  // Reset to page 1 and re-render
  _getState(tableId).page = 1;
  _renderTable(tableId);
}}

// ─────────────────────────────────────────────────────────────────────────────
// Tab switching
// ─────────────────────────────────────────────────────────────────────────────
function showTab(id) {{
  document.querySelectorAll('.panel').forEach(p => p.classList.remove('active'));
  document.querySelectorAll('.nav-item').forEach(n => n.classList.remove('active'));
  const panel = document.getElementById('panel-' + id);
  const nav   = document.getElementById('nav-' + id);
  if (panel) panel.classList.add('active');
  if (nav)   {{ nav.classList.add('active'); nav.scrollIntoView({{ block:'nearest' }}); }}
  // Initialise pagination for this tab on first visit
  if (id !== 'summary') _renderTable(id);
  const search = document.getElementById('search-' + id);
  if (search) setTimeout(() => search.focus(), 50);
}}

// ─────────────────────────────────────────────────────────────────────────────
// Boot: initialise pagination for ALL tables on load
// ─────────────────────────────────────────────────────────────────────────────
document.addEventListener('DOMContentLoaded', () => {{
  document.querySelectorAll('.panel[id^="panel-"]').forEach(panel => {{
    const id = panel.id.replace('panel-', '');
    if (id !== 'summary') _renderTable(id);
  }});
}});

// ─────────────────────────────────────────────────────────────────────────────
// Keyboard shortcuts
// ─────────────────────────────────────────────────────────────────────────────
document.addEventListener('keydown', e => {{
  if (e.target.tagName === 'INPUT') return;
  const activePanel = document.querySelector('.panel.active');
  if (!activePanel) return;
  const id = activePanel.id.replace('panel-', '');
  if (id === 'summary') return;

  if (e.key === '/') {{
    e.preventDefault();
    document.getElementById('search-' + id)?.focus();
  }} else if (e.key === 'ArrowRight' || e.key === ']') {{
    goPage(id, 'next');
  }} else if (e.key === 'ArrowLeft' || e.key === '[') {{
    goPage(id, 'prev');
  }}
}});
</script>
</body>
</html>"""


# ---------------------------------------------------------------------------
# Plain-text cell value (for Excel)
# ---------------------------------------------------------------------------

def _cell_text(value: Any, fmt: str) -> Any:
    """Return a plain-text / native Python value suitable for an Excel cell."""
    if value is None or value == "":
        return ""

    if fmt in ("plain", "mono", "path", "trunc"):
        s = str(value)
        return s[:200] if fmt == "trunc" else s

    if fmt == "short_mono":
        return str(value)[:8]

    if fmt == "badge_bool":
        return "Yes" if (value is True or str(value).lower() in ("true", "1", "yes")) else "No"

    if fmt in ("badge_state", "badge_type", "badge_lang"):
        return str(value)

    if fmt == "count":
        return len(value) if isinstance(value, list) else 0

    if fmt == "first_email":
        if isinstance(value, list) and value:
            return value[0].get("value", "")
        return ""

    if fmt == "list_vals":
        if isinstance(value, list) and value:
            return ", ".join(str(v.get("value", v)) for v in value)
        return ""

    if fmt == "schedule":
        if isinstance(value, dict):
            cron = value.get("quartz_cron_expression", "")
            tz   = value.get("timezone_id", "")
            return f"{cron}  ({tz})" if cron else "Manual"
        return "Manual"

    if fmt == "epoch_ms":
        try:
            ts = int(value) / 1000
            return datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M")
        except Exception:
            return str(value)

    if fmt == "iso_ts":
        try:
            return str(value)[:16].replace("T", " ")
        except Exception:
            return str(value)

    if fmt == "url_link":
        return str(value)

    if fmt == "kv_dns":
        return value.get("dns_name", "") if isinstance(value, dict) else ""

    return str(value)


# ---------------------------------------------------------------------------
# Excel Report Generator
# ---------------------------------------------------------------------------

_EXCEL_HEADER_BG  = "1E3A5F"   # deep navy
_EXCEL_SECTION_BG = "334155"   # slate
_EXCEL_DB_RED     = "FF3621"   # Databricks brand red
_EXCEL_ALT_ROW    = "F1F5F9"   # very light gray

_SUMMARY_CARD_KEYS = [
    "users", "groups", "service_principals",
    "notebooks", "workspace_files",
    "jobs", "clusters", "instance_pools", "cluster_policies",
    "sql_warehouses", "dlt_pipelines",
    "lakeview_dashboards", "genie_spaces",
    "secret_scopes", "repos", "serving_endpoints",
]


def _render_excel(workspace_url: str, data: Dict[str, Any],
                  errors: List[str], generated_at: str, output_path: str):
    """Generate an Excel workbook: Summary sheet + one sheet per component."""
    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("  ⚠  openpyxl not installed – skipping Excel (pip install openpyxl)")
        return

    def _fill(hex_color: str) -> PatternFill:
        return PatternFill("solid", fgColor=hex_color)

    def _font(bold=False, color="000000", size=10, italic=False) -> Font:
        return Font(bold=bold, color=color, size=size, italic=italic,
                    name="Calibri")

    def _align(h="left", v="center", wrap=False) -> Alignment:
        return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

    thin = Side(style="thin", color="CBD5E1")
    box_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    counts = _build_summary_data(data)
    wb = openpyxl.Workbook()

    # ── Summary sheet ─────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False

    # Title banner
    last_col = get_column_letter(3)
    ws.merge_cells(f"A1:{last_col}1")
    c = ws["A1"]
    c.value = "Databricks Workspace Inventory"
    c.font  = _font(bold=True, color="FFFFFF", size=16)
    c.fill  = _fill(_EXCEL_DB_RED)
    c.alignment = _align("center")
    ws.row_dimensions[1].height = 36

    # Sub-header
    ws.merge_cells(f"A2:{last_col}2")
    c = ws["A2"]
    c.value = f"Workspace: {workspace_url}   |   Generated: {generated_at}"
    c.font  = _font(italic=True, color="475569", size=9)
    c.fill  = _fill("FFF5F5")
    c.alignment = _align("center")
    ws.row_dimensions[2].height = 16

    ws.row_dimensions[3].height = 6  # spacer

    # Column headers
    for col_idx, header in enumerate(["Component", "Count", "Excel Sheet"], 1):
        c = ws.cell(row=4, column=col_idx, value=header)
        c.font   = _font(bold=True, color="FFFFFF", size=10)
        c.fill   = _fill(_EXCEL_HEADER_BG)
        c.alignment = _align("center")
        c.border = box_border
    ws.row_dimensions[4].height = 22

    # Data rows
    total = 0
    for row_idx, key in enumerate(_SUMMARY_CARD_KEYS, 5):
        count      = counts.get(key, 0)
        total     += count
        label      = _LABELS.get(key, key.replace("_", " ").title())
        sheet_name = re.sub(r'[\\/?*\[\]:]', '-', label)[:31]
        bg = _EXCEL_ALT_ROW if row_idx % 2 == 0 else "FFFFFF"

        for col_idx, val in enumerate([label, count, sheet_name], 1):
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.fill   = _fill(bg)
            c.border = box_border
            c.font   = _font(bold=(col_idx == 2), size=10)
            c.alignment = _align("center" if col_idx == 2 else "left")

    # Total row
    total_row = len(_SUMMARY_CARD_KEYS) + 5
    c = ws.cell(row=total_row, column=1, value="TOTAL")
    c.font   = _font(bold=True, size=10, color="FFFFFF")
    c.fill   = _fill(_EXCEL_DB_RED)
    c.border = box_border
    c.alignment = _align("center")

    c2 = ws.cell(row=total_row, column=2, value=total)
    c2.font   = _font(bold=True, size=11, color="FFFFFF")
    c2.fill   = _fill(_EXCEL_DB_RED)
    c2.border = box_border
    c2.alignment = _align("center")

    c3 = ws.cell(row=total_row, column=3, value="")
    c3.fill   = _fill(_EXCEL_DB_RED)
    c3.border = box_border

    ws.row_dimensions[total_row].height = 22

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 28
    ws.freeze_panes = "A5"

    # ── Component sheets ──────────────────────────────────────────────────
    for key in _SUMMARY_CARD_KEYS:
        items = data.get(key)
        if items is None:
            if key == "notebooks":
                items = [x for x in data.get("workspace_items", [])
                         if x.get("object_type") == "NOTEBOOK"]
            elif key == "workspace_files":
                items = [x for x in data.get("workspace_items", [])
                         if x.get("object_type") == "FILE"]
            else:
                items = []

        cols       = _COLUMNS.get(key, [("path", "Path", "plain")])
        label      = _LABELS.get(key, key.replace("_", " ").title())
        # Excel sheet names cannot contain: \ / ? * [ ] :
        sheet_name = re.sub(r'[\\/?*\[\]:]', '-', label)[:31]
        n_cols     = len(cols)

        ws2 = wb.create_sheet(title=sheet_name)
        ws2.sheet_view.showGridLines = False

        # Sheet title row
        ws2.merge_cells(f"A1:{get_column_letter(n_cols)}1")
        c = ws2["A1"]
        c.value     = f"{label}  —  {len(items):,} items"
        c.font      = _font(bold=True, color="FFFFFF", size=12)
        c.fill      = _fill(_EXCEL_HEADER_BG)
        c.alignment = _align("left")
        ws2.row_dimensions[1].height = 26

        # Sub-info row
        ws2.merge_cells(f"A2:{get_column_letter(n_cols)}2")
        c = ws2["A2"]
        c.value     = f"Workspace: {workspace_url}   |   Generated: {generated_at}"
        c.font      = _font(italic=True, color="64748B", size=9)
        c.fill      = _fill("F8FAFC")
        c.alignment = _align("left")
        ws2.row_dimensions[2].height = 14

        ws2.row_dimensions[3].height = 4  # spacer

        # Column header row
        for col_idx, (_, col_label, _) in enumerate(cols, 1):
            c = ws2.cell(row=4, column=col_idx, value=col_label)
            c.font      = _font(bold=True, color="FFFFFF", size=10)
            c.fill      = _fill(_EXCEL_SECTION_BG)
            c.alignment = _align("center")
            c.border    = box_border
        ws2.row_dimensions[4].height = 20
        ws2.freeze_panes = "A5"

        # Data rows
        col_widths = [len(col_label) for (_, col_label, _) in cols]
        for row_idx, item in enumerate(items, 5):
            bg = _EXCEL_ALT_ROW if row_idx % 2 == 0 else "FFFFFF"
            for col_idx, (key_path, _, fmt) in enumerate(cols, 1):
                raw  = _deep_get(item, key_path)
                text = _cell_text(raw, fmt)
                c    = ws2.cell(row=row_idx, column=col_idx, value=text)
                c.fill      = _fill(bg)
                c.font      = _font(size=10)
                c.alignment = _align("left", "top")
                c.border    = box_border
                if row_idx <= 104:  # sample first 100 rows for width
                    col_widths[col_idx - 1] = max(
                        col_widths[col_idx - 1],
                        min(len(str(text)) if text else 0, 60)
                    )

        # Apply column widths
        for col_idx, width in enumerate(col_widths, 1):
            ws2.column_dimensions[get_column_letter(col_idx)].width = min(width + 4, 68)

    # ── Warnings sheet ────────────────────────────────────────────────────
    if errors:
        ws_err = wb.create_sheet(title="Warnings")
        ws_err.sheet_view.showGridLines = False
        ws_err.merge_cells("A1:B1")
        c = ws_err["A1"]
        c.value     = "Fetch Warnings"
        c.font      = _font(bold=True, color="B91C1C", size=12)
        c.fill      = _fill("FEF2F2")
        c.alignment = _align("left")
        ws_err.row_dimensions[1].height = 24
        for row_idx, err in enumerate(errors, 2):
            ws_err.cell(row=row_idx, column=1, value=err).font = _font(size=10)
        ws_err.column_dimensions["A"].width = 90

    wb.save(output_path)
    print(f"  Excel report → {os.path.abspath(output_path)}\n")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    p = argparse.ArgumentParser(
        description="Generate Databricks workspace inventory HTML + Excel reports.")
    p.add_argument("--workspace-url", required=True, help="Databricks workspace URL")
    p.add_argument("--token",          required=True, help="Personal Access Token")
    p.add_argument("--output", default="",
                   help="Output HTML file path (auto-generated if omitted)")
    p.add_argument("--excel-output", default="",
                   help="Output Excel (.xlsx) file path (auto-generated if omitted)")
    p.add_argument("--no-excel", action="store_true",
                   help="Skip Excel report generation")
    p.add_argument("--max-scim", type=int, default=0,
                   help="Max items to fetch for Users/Groups/SPs "
                        "(0 = unlimited; use e.g. 2000 for large workspaces)")
    p.add_argument("--max-workspace-items", type=int, default=0,
                   help="Max workspace notebook/file items to scan "
                        "(0 = unlimited; use e.g. 5000 for large workspaces)")
    p.add_argument("--max-ws-api-calls", type=int, default=300,
                   help="Max workspace/list API calls during recursive scan "
                        "(default: 300; prevents long scans on deep source trees)")
    p.add_argument("--no-ssl-verification", action="store_true",
                   help="Disable SSL certificate verification")
    p.add_argument("--verbose", "-v", action="store_true",
                   help="Print each API call, pagination page, and timing details")
    args = p.parse_args()

    if args.verbose:
        vlog.enable()

    workspace_url = args.workspace_url.rstrip("/")
    verify_ssl    = not args.no_ssl_verification
    hostname      = urlparse(workspace_url).hostname or workspace_url
    ts_file       = datetime.now().strftime("%Y%m%d%H%M")
    ts_display    = datetime.now().strftime("%Y-%m-%d %H:%M")

    base_name    = f"inventory_{hostname.split('.')[0]}_{ts_file}"
    html_output  = args.output       or f"{base_name}.html"
    excel_output = args.excel_output or f"{base_name}.xlsx"

    scim_note = f" (capped at {args.max_scim} per type)" if args.max_scim else ""

    print(f"\n  Databricks Workspace Inventory")
    print(f"  {'─' * 54}")
    print(f"  Workspace  : {workspace_url}")
    print(f"  HTML       : {html_output}")
    if not args.no_excel:
        print(f"  Excel      : {excel_output}")
    if args.max_scim:
        print(f"  SCIM limit : {args.max_scim} per type")
    if args.max_workspace_items:
        print(f"  WS items   : {args.max_workspace_items} max")
    print(f"  WS API cap : {args.max_ws_api_calls} list calls max")
    print(f"  {'─' * 54}\n")

    client    = DatabricksClient(workspace_url, args.token, verify_ssl=verify_ssl)
    inventory = WorkspaceInventory(client, max_scim=args.max_scim,
                                   max_workspace_items=args.max_workspace_items,
                                   max_ws_api_calls=args.max_ws_api_calls)

    t0 = time.time()
    data = inventory.fetch_all()
    elapsed = time.time() - t0

    counts = _build_summary_data(data)
    total  = sum(counts.values())

    print(f"\n  {'─' * 54}")
    print(f"  Inventory complete in {elapsed:.1f}s — {total:,} total resources{scim_note}\n")
    for key, count in counts.items():
        icon = _ICONS.get(key, ("📦",))[0]
        label = _LABELS.get(key, key)
        print(f"    {icon}  {label:<30} {count:>6,}")
    if inventory.errors:
        print(f"\n  ⚠  Warnings ({len(inventory.errors)}):")
        for e in inventory.errors:
            print(f"     • {e}")
    print()

    html = _render_html(workspace_url, data, inventory.errors, ts_display)
    with open(html_output, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"  HTML report → {os.path.abspath(html_output)}\n")

    if not args.no_excel:
        _render_excel(workspace_url, data, inventory.errors, ts_display, excel_output)


if __name__ == "__main__":
    main()
